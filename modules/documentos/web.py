from __future__ import annotations

from fastapi import APIRouter, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates

from modules.documentos.services.dashboard import (
    dashboard_snapshot,
    export_documents,
)

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

router = APIRouter()
templates = Jinja2Templates(directory="web/templates")


@router.get("/documentos", response_class=HTMLResponse)
def documentos_dashboard(
    request: Request,
    page: int = 1,
    per_page: int = 25,
    q: str = "",
    status: str = "",
    carrier_cnpj: str = "",
):
    if not request.session.get("ssw_session_id"):
        return RedirectResponse("/login", status_code=303)

    snapshot = dashboard_snapshot(
        page=page,
        per_page=per_page,
        search=q,
        alert_type=status,
        carrier_cnpj=carrier_cnpj,
    )

    return templates.TemplateResponse(
        "documentos_dashboard_v2.html",
        {
            "request": request,
            "dashboard": snapshot,
        },
    )

@router.get("/documentos/export")
def documentos_export(
    request: Request,
    q: str = "",
    status: str = "",
    carrier_cnpj: str = "",
):
    if not request.session.get("ssw_session_id"):
        return RedirectResponse("/login", status_code=303)

    rows = export_documents(
        search=q,
        alert_type=status,
        carrier_cnpj=carrier_cnpj,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Acompanhamento"

    headers = [
        "Situação",
        "CNPJ Transportadora",
        "Nota Fiscal",
        "Série NF",
        "CTRC",
        "Destinatário",
        "CNPJ Destinatário",
        "Número Transporte",
        "Data Entrega",
        "Prazo",
        "Dias Restantes",
        "Status Documental",
        "Motivo Pendência",
        "Usuário Pendência",
        "Data Pendência",
        "Quantidade Scans",
        "Último Scan",
        "Unidades Scan",
        "Parceiros",
        "Pacote Arquivo",
        "Capa Remessa",
        "Mapa Expedição",
        "Mapa Recepção",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    status_labels = {
        "VENCIDO": "Vencido",
        "ALERTA_5": "Até 5 dias",
        "ALERTA_10": "Até 10 dias",
        "ALERTA_20": "Até 20 dias",
        "SEM_OC_01": "Sem data de entrega",
        "NORMAL": "Regular",
    }

    def excel_value(value):
        if value is None:
            return ""

        if isinstance(value, (date, datetime)):
            return value

        return str(value)

    for row in rows:
        details = row["details"]

        worksheet.append(
            [
                status_labels.get(
                    row.get("alert_type"),
                    row.get("alert_type") or "",
                ),
                excel_value(details.get("carrier_cnpj")),
                excel_value(details.get("invoice_number")),
                excel_value(details.get("invoice_series")),
                excel_value(details.get("ctrc")),
                excel_value(details.get("recipient_name")),
                excel_value(details.get("recipient_cnpj")),
                excel_value(details.get("transport_number")),
                excel_value(row.get("delivery_date")),
                excel_value(row.get("deadline_date")),
                row.get("days_remaining"),
                (
                    "Recusado / com pendência"
                    if details.get("has_pending")
                    else "Sem pendência"
                ),
                excel_value(details.get("pending_reason")),
                excel_value(details.get("pending_user")),
                excel_value(details.get("pending_at")),
                details.get("scan_count") or 0,
                excel_value(details.get("last_scan_at")),
                excel_value(details.get("scan_units")),
                excel_value(details.get("scan_partners")),
                excel_value(details.get("archive_package_number")),
                excel_value(details.get("remittance_cover_number")),
                excel_value(details.get("expedition_map")),
                excel_value(details.get("reception_map")),
            ]
        )

    # Congela cabeçalho.
    worksheet.freeze_panes = "A2"

    # Filtro nativo do Excel.
    worksheet.auto_filter.ref = worksheet.dimensions

    # Formatação das datas.
    for row_number in range(2, worksheet.max_row + 1):
        for column_number in (9, 10, 15, 17):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"

            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"

    # Largura automática com limite.
    for column_cells in worksheet.columns:
        column_index = column_cells[0].column
        column_letter = get_column_letter(column_index)

        max_length = 0

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            max_length = max(
                max_length,
                len(str(value)),
            )

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            45,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        "acompanhamento_documental_"
        f"{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )

@router.get(
    "/documentos/metricas",
    response_class=HTMLResponse,
)
def documentos_metricas(
    request: Request,
    days: int = 30,
):
    from modules.documentos.services.metrics import (
        metrics_snapshot,
    )

    days = max(7, min(days, 365))

    dashboard = metrics_snapshot(days)

    return templates.TemplateResponse(
        "documentos_metricas.html",
        {
            "request": request,
            "dashboard": dashboard,
            "days": days,
        },
    )