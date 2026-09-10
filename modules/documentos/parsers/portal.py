from __future__ import annotations

from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from modules.documentos.schemas import PortalDocument
from modules.documentos.services.normalizer import (
    clean_text,
    normalize_cnpj,
    normalize_ctrc,
    normalize_identifier,
    optional_text,
    parse_date,
)


REQUIRED_COLUMNS = {
    "CNPJ Transportador",
    "CTRC",
    "Número",
    "Transporte",
    "Data Emissão",
}


def _pending_datetime(value: object):
    # O Excel mantém a fração do dia no valor serial.
    from datetime import datetime, timedelta

    if isinstance(value, datetime):
        return value
    if isinstance(value, (int, float)):
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    parsed = parse_date(value)
    return datetime.combine(parsed, datetime.min.time()) if parsed else None


def iter_portal(file_path: Path, report_type: str) -> Iterator[PortalDocument]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.active

    # Alguns arquivos do portal declaram incorretamente dimension ref="A1".
    if worksheet.max_row == 1 and worksheet.max_column == 1:
        worksheet.reset_dimensions()

    try:
        rows = worksheet.iter_rows(values_only=True)
        header = [clean_text(value) for value in next(rows)]
        missing = REQUIRED_COLUMNS - set(header)
        if missing:
            raise ValueError(
                "Layout do portal inválido. Colunas ausentes: "
                + ", ".join(sorted(missing))
            )

        for values in rows:
            padded = list(values) + [None] * max(0, len(header) - len(values))
            record = dict(zip(header, padded))
            invoice_number = normalize_identifier(record.get("Número"))
            if not invoice_number:
                continue

            ctrc_raw = clean_text(record.get("CTRC"))
            yield PortalDocument(
                report_type=report_type,
                carrier_cnpj=normalize_cnpj(record.get("CNPJ Transportador")),
                ctrc=normalize_ctrc(ctrc_raw),
                ctrc_raw=ctrc_raw,
                invoice_number=invoice_number,
                invoice_series=normalize_identifier(record.get("Série")) or None,
                transport_number=normalize_identifier(record.get("Transporte")) or None,
                issue_date=parse_date(record.get("Data Emissão")),
                recipient_name=optional_text(record.get("Nome Destinatário")),
                recipient_cnpj=normalize_cnpj(record.get("CNPJ Destinatário")),
                recipient_state=optional_text(record.get("Estado Destinatário")),
                pending_at=_pending_datetime(record.get("Data Pendência")),
                pending_user=optional_text(record.get("Usuário Pendência")),
                pending_reason=optional_text(record.get("Pendência")),
            )
    finally:
        workbook.close()

