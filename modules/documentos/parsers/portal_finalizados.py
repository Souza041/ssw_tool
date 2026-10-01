from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from modules.documentos.services.normalizer import (
    clean_text,
    normalize_cnpj,
    normalize_identifier,
    optional_text,
    parse_date,
)


REQUIRED_COLUMNS = {
    "CNPJ Transportador",
    "CTRC",
    "Número",
    "Série",
    "Data Emissão",
    "Finalizado",
}


@dataclass(frozen=True)
class PortalFinalizado:
    carrier_cnpj: str
    ctrc: str | None
    invoice_number: str
    invoice_series: str | None
    issue_date: date | None
    finalized_at: date
    transport_number: str | None
    recipient_name: str | None
    recipient_cnpj: str | None
    recipient_state: str | None


def iter_portal_finalizados(
    file_path: Path,
) -> Iterator[PortalFinalizado]:

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    if worksheet.max_row == 1 and worksheet.max_column == 1:
        worksheet.reset_dimensions()

    try:
        rows = worksheet.iter_rows(values_only=True)

        try:
            first_row = next(rows)
        except StopIteration:
            return

        header = [
            clean_text(value)
            for value in first_row
        ]

        missing = REQUIRED_COLUMNS - set(header)

        if missing:
            raise ValueError(
                "Layout de Finalizados inválido. Colunas ausentes: "
                + ", ".join(sorted(missing))
            )

        for values in rows:
            padded = list(values) + [None] * max(
                0,
                len(header) - len(values),
            )

            record = dict(zip(header, padded))

            invoice_number = normalize_identifier(
                record.get("Número")
            )

            finalized_at = parse_date(
                record.get("Finalizado")
            )

            carrier_cnpj = normalize_cnpj(
                record.get("CNPJ Transportador")
            )

            if not invoice_number:
                continue

            if not carrier_cnpj:
                continue

            if finalized_at is None:
                continue

            yield PortalFinalizado(
                carrier_cnpj=carrier_cnpj,
                ctrc=normalize_identifier(
                    record.get("CTRC")
                ) or None,
                invoice_number=invoice_number,
                invoice_series=normalize_identifier(
                    record.get("Série")
                ) or None,
                issue_date=parse_date(
                    record.get("Data Emissão")
                ),
                finalized_at=finalized_at,
                transport_number=normalize_identifier(
                    record.get("Transporte")
                ) or None,
                recipient_name=optional_text(
                    record.get("Nome Destinatário")
                ),
                recipient_cnpj=normalize_cnpj(
                    record.get("CNPJ Destinatário")
                ),
                recipient_state=optional_text(
                    record.get("Estado")
                ),
            )

    finally:
        workbook.close()