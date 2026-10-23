import argparse
import json
import re
from collections import Counter

from openpyxl import load_workbook

from .simular_consolidacao import consultar_com_retry


def normalizar_numero(valor):
    valor = str(valor).replace("\u00a0", "").strip()

    if not re.fullmatch(r"\d+(?:\.0+)?", valor):
        return None

    numero = str(int(valor.split(".")[0]))

    if numero == "0" or len(numero) > 9:
        return None

    return numero


def ler_ctes(caminho):
    workbook = load_workbook(
        caminho,
        read_only=True,
        data_only=True,
    )

    try:
        aba = workbook["Base termo"]

        if str(aba["D1"].value).strip().upper() != "# CT-E":
            raise ValueError("Cabeçalho inesperado na coluna D.")

        ctes = Counter()
        invalidos = 0

        for (valor,) in aba.iter_rows(
            min_row=2,
            min_col=4,
            max_col=4,
            values_only=True,
        ):
            if valor is None or not str(valor).strip():
                continue

            numero = normalizar_numero(valor)

            if numero is None:
                invalidos += 1
            else:
                ctes[numero] += 1

        return ctes, invalidos

    finally:
        workbook.close()


def diagnosticar(caminho):
    ctes, invalidos = ler_ctes(caminho)
    numeros = sorted(ctes)
    candidatos = {numero: set() for numero in numeros}

    for inicio in range(0, len(numeros), 100):
        lote = numeros[inicio:inicio + 100]
        marcadores = ",".join(["%s"] * len(lote))

        registros = consultar_com_retry(f"""
            SELECT DISTINCT
                r.cte_numero,
                r.cte_serie,
                r.ctrc_chave
            FROM lg_registros_importados AS r
            INNER JOIN lg_importacoes AS i
                ON i.id = r.importacao_id
            WHERE i.status = 'concluida'
              AND i.relatorio IN ('OP435', 'OP455')
              AND r.cte_numero IN ({marcadores})
              AND r.ctrc_chave IS NOT NULL
              AND r.ctrc_chave <> ''
        """, tuple(lote))

        for registro in registros:
            candidatos[registro["cte_numero"]].add((
                registro["cte_serie"],
                registro["ctrc_chave"],
            ))

        print(
            f"Consultados {min(inicio + 100, len(numeros))}"
            f"/{len(numeros)} CT-e.",
            flush=True,
        )

    unicos = {
        numero for numero, opcoes in candidatos.items()
        if len(opcoes) == 1
    }
    ausentes = {
        numero for numero, opcoes in candidatos.items()
        if not opcoes
    }
    ambiguos = {
        numero for numero, opcoes in candidatos.items()
        if len(opcoes) > 1
    }

    print(json.dumps({
        "linhas_com_cte_valido": sum(ctes.values()),
        "ctes_distintos": len(ctes),
        "linhas_com_cte_invalido": invalidos,
        "ctes_com_correspondencia_unica": len(unicos),
        "ctes_sem_correspondencia": len(ausentes),
        "ctes_com_correspondencia_ambigua": len(ambiguos),
        "amostras_sem_correspondencia": sorted(ausentes)[:10],
        "amostras_ambiguas": [
            {
                "cte": numero,
                "candidatos": [
                    {"serie": serie, "ctrc": ctrc}
                    for serie, ctrc in sorted(
                        candidatos[numero],
                        key=lambda item: (item[0] or "", item[1]),
                    )
                ],
            }
            for numero in sorted(ambiguos)[:10]
        ],
    }, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(
        description="Confere correspondências dos CT-e da Base termo."
    )
    parser.add_argument("--planilha", required=True)
    args = parser.parse_args()
    diagnosticar(args.planilha)


if __name__ == "__main__":
    main()
