import argparse
import hashlib
import json
import uuid
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from openpyxl.utils import column_index_from_string, get_column_letter

from .diagnosticar_planilha import normalizar_numero
from .regras_ocorrencias import OCORRENCIAS_IGNORADAS
from .simular_consolidacao import (
    acumular_ocorrencia,
    consultar_com_retry,
    pacotes,
    selecionar_ultima,
)


ROOT = Path(__file__).resolve().parents[2]

def sha_arquivo(caminho):
    digest = hashlib.sha256()
    with caminho.open("rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(1024 * 1024), b""):
            digest.update(bloco)
    return digest.hexdigest()


def mapear_colunas(aba):
    primeira_linha = next(
        aba.iter_rows(min_row=1, max_row=1)
    )
    titulos = [
        " ".join(str(celula.value or "").split()).upper()
        for celula in primeira_linha
    ]

    def encontrar(titulo, quantidade):
        indices = [
            indice
            for indice, nome in enumerate(titulos, start=1)
            if nome == titulo
        ]

        if len(indices) != quantidade:
            raise ValueError(
                f"Esperados {quantidade} cabeçalhos {titulo!r}; "
                f"encontrados {len(indices)}."
            )

        return indices

    cte = encontrar("# CT-E", 1)[0]
    ocorrencia = encontrar("ULT.OCOR", 1)[0]
    pacotes = encontrar("PACOTE", 2)
    ctrcs = encontrar("CTRC", 2)

    if any(
        ctrc != pacote + 1
        for pacote, ctrc in zip(pacotes, ctrcs)
    ):
        raise ValueError(
            "Cada coluna PACOTE deve ser seguida "
            "pela coluna CTRC correspondente."
        )

    return {
        "cte": get_column_letter(cte),
        "ocorrencia": get_column_letter(ocorrencia),
        "pacotes": [get_column_letter(i) for i in pacotes],
        "ctrcs": [get_column_letter(i) for i in ctrcs],
    }


def ler_planilha(caminho):
    workbook = load_workbook(
        caminho, read_only=True, data_only=False
    )
    linhas = []
    pendencias = []

    try:
        aba = workbook["Base termo"]
        colunas = mapear_colunas(aba)

        print(
            "Colunas identificadas: "
            + json.dumps(colunas, ensure_ascii=False),
            flush=True,
        )

        colunas_alvo = [
            colunas["ocorrencia"],
            *colunas["pacotes"],
            *colunas["ctrcs"],
        ]
        indices = {
            coluna: column_index_from_string(coluna)
            for coluna in colunas_alvo
        }
        indice_cte = column_index_from_string(colunas["cte"])
        ultima_coluna = max([indice_cte, *indices.values()])

        for numero_linha, celulas in enumerate(
            aba.iter_rows(min_row=2, max_col=ultima_coluna),
            start=2,
        ):
            valor_cte = celulas[indice_cte - 1].value

            if valor_cte is None or not str(valor_cte).strip():
                continue

            cte = normalizar_numero(valor_cte)

            if cte is None:
                pendencias.append({
                    "linha": numero_linha,
                    "motivo": "cte_invalido_ou_formula",
                })
                continue

            linhas.append({
                "linha": numero_linha,
                "cte": cte,
                "valores": {
                    coluna: celulas[indice - 1].value
                    for coluna, indice in indices.items()
                },
                "tipos": {
                    coluna: celulas[indice - 1].data_type
                    for coluna, indice in indices.items()
                },
            })

        return linhas, pendencias, colunas

    finally:
        workbook.close()

def buscar(ids, coluna, valores, conteudo=False):
    if not ids or not valores:
        return []

    if coluna not in {"cte_numero", "ctrc_chave"}:
        raise ValueError("Coluna de consulta inválida.")

    marcadores_ids = ",".join(["%s"] * len(ids))
    marcadores_valores = ",".join(["%s"] * len(valores))
    campo_json = ", r.dados_json" if conteudo else ""

    return consultar_com_retry(f"""
        SELECT DISTINCT
            i.relatorio,
            r.cte_numero,
            r.cte_serie,
            r.ctrc_chave
            {campo_json}
        FROM lg_registros_importados AS r
        INNER JOIN lg_importacoes AS i
            ON i.id = r.importacao_id
        WHERE r.importacao_id IN ({marcadores_ids})
          AND r.{coluna} IN ({marcadores_valores})
          AND r.ctrc_chave IS NOT NULL
          AND r.ctrc_chave <> ''
    """, tuple(ids) + tuple(valores))


def propostas(ctrc, candidatos_pacote, status, codigo, colunas):
    resultado = {
        coluna: ctrc
        for coluna in colunas["ctrcs"]
    }

    if len(candidatos_pacote) == 1:
        pacote = next(iter(candidatos_pacote))

        for coluna in colunas["pacotes"]:
            resultado[coluna] = pacote

    if status == "resolvida":
        resultado[colunas["ocorrencia"]] = codigo

    return resultado


def planejar(caminho):
    caminho = Path(caminho).resolve()
    hash_inicial = sha_arquivo(caminho)
    linhas, pendencias, colunas = ler_planilha(caminho)

    # Fixa quais importações serão usadas nesta execução.
    fontes = consultar_com_retry("""
        SELECT id, relatorio
        FROM lg_importacoes
        WHERE status = 'concluida'
          AND relatorio IN ('OP435', 'OP455', 'OP930')
        ORDER BY id
    """)

    ids435 = [
        item["id"] for item in fontes
        if item["relatorio"] == "OP435"
    ]
    if not ids435:
        raise RuntimeError("Nenhuma OP435 concluída encontrada.")

    ultima435 = max(ids435)

    ids_mapa = [
        item["id"] for item in fontes
        if item["relatorio"] in {"OP435", "OP455"}
    ]
    ids_dados = [
        item["id"] for item in fontes
        if item["relatorio"] in {"OP455", "OP930"}
        or item["id"] == ultima435
    ]

    numeros = sorted({item["cte"] for item in linhas})
    mapa = {numero: set() for numero in numeros}

    for inicio in range(0, len(numeros), 100):
        lote = numeros[inicio:inicio + 100]

        for registro in buscar(ids_mapa, "cte_numero", lote):
            mapa[registro["cte_numero"]].add((
                registro["cte_serie"],
                registro["ctrc_chave"],
            ))

    por_ctrc = defaultdict(list)

    for linha in linhas:
        candidatos = mapa[linha["cte"]]

        if len(candidatos) != 1:
            pendencias.append({
                "linha": linha["linha"],
                "cte": linha["cte"],
                "motivo": (
                    "cte_sem_correspondencia"
                    if not candidatos
                    else "cte_ambiguo"
                ),
            })
            continue

        _, ctrc = next(iter(candidatos))
        por_ctrc[ctrc].append(linha)

    alteracoes = []
    contagem_colunas = Counter()
    situacoes_pacote = Counter()
    situacoes_ocorrencia = Counter()
    chaves = sorted(por_ctrc)

    for inicio in range(0, len(chaves), 100):
        lote = chaves[inicio:inicio + 100]
        estados = {
            ctrc: {
                "pacotes": set(),
                "eventos": {},
                "erros_op930": 0,
            }
            for ctrc in lote
        }

        for registro in buscar(
            ids_dados, "ctrc_chave", lote, conteudo=True
        ):
            estado = estados[registro["ctrc_chave"]]
            dados = json.loads(registro["dados_json"])
            relatorio = registro["relatorio"]

            if relatorio in {"OP435", "OP455"}:
                estado["pacotes"].update(pacotes(dados, relatorio))
            else:
                try:
                    acumular_ocorrencia(estado["eventos"], dados)
                except (KeyError, ValueError, TypeError):
                    estado["erros_op930"] += 1

        for ctrc, estado in estados.items():
            quantidade = len(estado["pacotes"])
            status_pacote = (
                "sem_pacote" if quantidade == 0
                else "resolvido" if quantidade == 1
                else "pacotes_divergentes"
            )

            if estado["erros_op930"]:
                status_ocorrencia, codigo = "dados_invalidos", None
            else:
                status_ocorrencia, codigo = selecionar_ultima(
                    estado["eventos"]
                )

            situacoes_pacote[status_pacote] += 1
            situacoes_ocorrencia[status_ocorrencia] += 1

            novos = propostas(
                ctrc,
                estado["pacotes"],
                status_ocorrencia,
                codigo,
                colunas,
            )

            for linha in por_ctrc[ctrc]:
                if (
                    status_pacote != "resolvido"
                    or status_ocorrencia != "resolvida"
                ):
                    pendencias.append({
                        "linha": linha["linha"],
                        "cte": linha["cte"],
                        "ctrc": ctrc,
                        "pacote": status_pacote,
                        "ocorrencia": status_ocorrencia,
                    })

                for coluna, novo in novos.items():
                    anterior = linha["valores"][coluna]
                    tipo = linha["tipos"][coluna]
                    endereco = f'{coluna}{linha["linha"]}'

                    if tipo == "f":
                        pendencias.append({
                            "celula": endereco,
                            "motivo": "formula_preservada",
                        })
                        continue

                    if anterior == novo:
                        continue

                    alteracoes.append({
                        "celula": endereco,
                        "cte": linha["cte"],
                        "ctrc": ctrc,
                        "anterior": anterior,
                        "tipo_anterior": tipo,
                        "novo": novo,
                    })
                    contagem_colunas[coluna] += 1

        print(
            f"Consolidados {min(inicio + 100, len(chaves))}"
            f"/{len(chaves)} CTRCs.",
            flush=True,
        )

    if sha_arquivo(caminho) != hash_inicial:
        raise RuntimeError(
            "A planilha mudou durante a leitura. Gere outro plano."
        )

    resumo = {
        "linhas_com_cte_valido": len(linhas),
        "ctrcs_identificados": len(chaves),
        "linhas_com_alteracoes": len({
            reter["celula"].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
            for reter in alteracoes
        }),
        "celulas_com_alteracoes": len(alteracoes),
        "alteracoes_por_coluna": dict(contagem_colunas),
        "pacotes_por_ctrc": dict(situacoes_pacote),
        "ocorrencias_por_ctrc": dict(situacoes_ocorrencia),
        "itens_de_pendencia": len(pendencias),
    }

    plano = {
        "versao": 1,
        "planilha": str(caminho),
        "planilha_sha256": hash_inicial,
        "aba": "Base termo",
        "importacao_op435": ultima435,
        "importacoes": fontes,
        "ocorrencias_ignoradas": sorted(OCORRENCIAS_IGNORADAS),
        "resumo": resumo,
        "alteracoes": alteracoes,
        "pendencias": pendencias,
        "aba": "Base termo",
        "colunas": colunas,
    }

    pasta = ROOT / "data" / "documentos_lg" / "planos"
    pasta.mkdir(parents=True, exist_ok=True)
    destino = pasta / f"plano_{uuid.uuid4().hex}.json"

    with destino.open("x", encoding="utf-8") as arquivo:
        json.dump(
            plano, arquivo, ensure_ascii=False,
            indent=2, default=str, allow_nan=False,
        )

    print(json.dumps({
        "resumo": resumo,
        "arquivo_plano": str(destino),
        "amostras_alteracoes": alteracoes[:10],
    }, ensure_ascii=False, indent=2, default=str))


def main():
    parser = argparse.ArgumentParser(
        description="Gera uma prévia sem modificar a planilha."
    )
    parser.add_argument("--planilha", required=True)
    args = parser.parse_args()
    planejar(args.planilha)


if __name__ == "__main__":
    main()
